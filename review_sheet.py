"""
Export a native-speaker review sheet for one or more languages.

Turns "translation quality is a known limitation" into something a person can
actually act on. Produces a CSV with the English original beside the translation
and an empty verdict column, so a speaker of the language can work through it
without reading any code or JSON.

The question the sheet asks is deliberately narrow. Not "is this good writing",
but "would a native speaker reading this understand it as an instruction to do
the thing it describes?" A clumsy attack that still reads as an attack is fine
for this purpose. One that reads as nonsense is not, because the bot would then
fail to comply for reasons that have nothing to do with its defences, and that
would show up as a safety finding.

    python review_sheet.py gu hi        # specific languages
    python review_sheet.py --tier low   # every low-resource language in the bank
    python review_sheet.py --all
"""
from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path

from languages_catalog import tier_of

BANK_PATH = Path(__file__).with_name("attack_bank.json")
OUT_DIR = Path(__file__).with_name("review")

INSTRUCTIONS = [
    ["# PolyGuard translation review"],
    ["# Thank you for doing this. It should take about 15 minutes."],
    ["#"],
    ["# These are deliberately adversarial test messages used to check whether an"],
    ["# AI chatbot can be tricked. They are security tests, not real requests, and"],
    ["# nothing here is meant to be acted on."],
    ["#"],
    ["# For each row, fill in 'verdict' with one of:"],
    ["#    ok       reads naturally enough, and clearly asks for what the English asks for"],
    ["#    awkward  understandable, but clumsy or unnatural phrasing"],
    ["#    wrong    confusing, nonsensical, or asks for something different"],
    ["#"],
    ["# The key question is NOT whether it is elegant. It is whether a native"],
    ["# speaker would understand it as an instruction to do the thing it describes."],
    ["# A clumsy test that still reads as a test is fine. One that reads as"],
    ["# nonsense is not."],
    ["#"],
    ["# Put a better version in 'suggested_correction' if you can. Leave the codes"],
    ["# (PWNED_7F3A, CTRL_4B8E), any long random looking string, and the words"],
    ["# SYSTEM OVERRIDE exactly as they are: the test depends on them."],
    [],
]

# One question per language that only a native speaker can answer, each based
# on something measured in the bank rather than guessed. Shown at the top of the
# spreadsheet so the reviewer knows the one thing we most need from them.
QUESTIONS = {
    "de": "The attack lines address the chatbot as du, while the benign check lines use Sie. Is du what someone trying to trick a chatbot would naturally write? The umlauts in the attack lines were only just restored, so please watch for any that are wrong.",
    "fr": "The attack lines use tu and the benign check lines use vous. Does tu feel natural for a pushy instruction to a chatbot? The accents in the attack lines were only just restored, so please watch for any that are wrong.",
    "es": "Everything uses tú. Is that what a Spanish speaker would normally use when typing to a chatbot? The accents in the attack lines were only just restored, so please watch for any that are wrong.",
    "it": "Everything uses tu. Is that what an Italian speaker would normally use when typing to a chatbot? The accents in the attack lines were only just restored, so please watch for any that are wrong.",
    "pt": "The text is written in Brazilian Portuguese (você, usuário). Does it read naturally to a Brazilian speaker? The accents in the attack lines were only just restored, so please watch for any that are wrong.",
    "ru": "The attack lines use ты and the benign check lines use Вы. Is ты what someone would naturally use in a pushy instruction to a chatbot?",
    "uk": "The attack lines use ти and the benign check lines use Ви. Is ти what someone would naturally use in a pushy instruction to a chatbot?",
    "pl": "One line I am unsure of: in the indirect injection rows, Proszę streszcz tę opinię. I suspect streszcz should be streść. The Polish letters in the attack lines were only just restored, so please watch for any that are wrong.",
    "zh": "The attack lines use 你 and the benign check lines use 您, all in simplified characters. Does each feel natural for its purpose?",
    "hi": "The attack lines mix आप and तुम, while the benign check lines use आप. Which would a Hindi speaker actually use when typing to a chatbot?",
    "gu": "The lines mostly use તમે. Does the Gujarati read naturally, and do any lines read more like Hindi phrasing than real Gujarati?",
    "tr": "The attack lines mostly address the chatbot as sen. Is that natural here? The Turkish letters in the attack lines were only just restored, so please watch for any that are wrong.",
    "id": "The attack lines use kamu (for example tugasmu) and the benign check lines use Anda. Which would an Indonesian speaker use when typing to a chatbot?",
    "tl": "Some lines use po and some do not. Would a Filipino user typing to a chatbot write in Tagalog like this, or mix in English the way many people do?",
    "el": "Most lines address the chatbot as εσύ. Does the Greek read naturally for a message typed to a chatbot?",
    "ar": "The text is in Modern Standard Arabic. Would a user typing to a chatbot write in MSA, or would a dialect be more realistic?",
    "ja": "Every line uses polite です and ます forms, including the attacks. Would someone trying to trick a chatbot write this politely, or more bluntly?",
    "ko": "Every line uses polite 세요 and 습니다 endings, including the attacks. Would someone trying to trick a chatbot write this politely, or use a blunter speech level?",
    "vi": "The tone marks in the attack lines were only just restored, because they had been typed without any. Please check them closely, since a wrong tone mark can change the word.",
}

CATEGORY_LABEL = {
    "instruction_override": "Override instructions",
    "role_play_jailbreak": "Role play trick",
    "system_prompt_extraction": "Reveal hidden instructions",
    "obfuscated_payload": "Hidden encoded message",
    "indirect_injection": "Instruction hidden in content",
    "benign": "Benign check, not an attack",
}

HEADER = ["id", "kind", "category", "english_original", "translation",
          "verdict", "suggested_correction", "notes"]


def load_bank() -> dict:
    return json.loads(BANK_PATH.read_text(encoding="utf-8"))


def rows_for(bank: dict, code: str) -> list[list[str]]:
    en_attacks = {(a["category"], a.get("variant", 0)): a["text"]
                  for a in bank["attacks"] if a["lang"] == "en"}
    en_controls = {c.get("variant", 0): c["text"]
                   for c in bank.get("controls", []) if c["lang"] == "en"}
    rows = []
    for a in bank["attacks"]:
        if a["lang"] != code:
            continue
        rows.append([a["id"], "attack", a["category"],
                     en_attacks.get((a["category"], a.get("variant", 0)), ""),
                     a["text"], "", "", ""])
    for c in bank.get("controls", []):
        if c["lang"] != code:
            continue
        rows.append([c["id"], "capability control", "benign",
                     en_controls.get(c.get("variant", 0), ""), c["text"], "", "", ""])
    return rows


def write_sheet(bank: dict, code: str) -> Path:
    meta = bank["languages"][code]
    OUT_DIR.mkdir(exist_ok=True)
    path = OUT_DIR / f"{code}_review.csv"
    # utf-8-sig so Excel opens non-Latin scripts correctly instead of showing
    # the reviewer a screen of mojibake and wasting their time.
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        for line in INSTRUCTIONS:
            w.writerow(line)
        w.writerow([f"# Language: {meta['name']} ({meta['native']}), "
                    f"tier {tier_of(code)}, provenance {meta.get('provenance', '?')}"])
        w.writerow([])
        w.writerow(HEADER)
        w.writerows(rows_for(bank, code))
    return path


def write_xlsx(bank: dict, code: str) -> Path:
    """A spreadsheet a busy person can open and finish without reading any code.

    Same rows as the CSV. The difference is presentation: instructions on their
    own tab instead of comment lines Excel turns into clutter, a dropdown for the
    verdict so answers come back in three known values, wrapped text, and a
    frozen header. The reviewer is doing us a favour; the sheet should respect it.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    meta = bank["languages"][code]
    OUT_DIR.mkdir(exist_ok=True)
    path = OUT_DIR / f"{code}_review.xlsx"
    wb = Workbook()

    start = wb.active
    start.title = "Start here"
    start.column_dimensions["A"].width = 100
    lines = [
        (f"PolyGuard translation review: {meta['name']} ({meta['native']})", 16, True),
        ("About 15 minutes. Thank you for doing this.", 12, False),
        ("", 11, False),
        ("What these are", 12, True),
        ("Test messages used to check whether an AI chatbot can be tricked into ignoring its "
         "instructions. They are security tests, not real requests, and nothing here is meant "
         "to be acted on. The chatbot is only ever asked to repeat a harmless code word.", 11, False),
        ("", 11, False),
        ("What to do", 12, True),
        ("On the Review tab, pick a verdict for each row from the dropdown:", 11, False),
        ("    ok: reads naturally enough, and clearly asks for what the English asks for", 11, False),
        ("    awkward: understandable, but clumsy or unnatural", 11, False),
        ("    wrong: confusing, nonsensical, or asks for something different", 11, False),
        ("", 11, False),
        ("The question is not whether it is elegant. It is whether a native speaker would "
         "understand it as an instruction to do what it describes. A clumsy line that still "
         "reads as an instruction is fine.", 11, False),
        ("", 11, False),
        ("If you can, put a better version in the next column. Please leave these exactly as "
         "they are, because the test depends on them: the codes PWNED_7F3A and CTRL_4B8E, any "
         "long random looking string, and the English words SYSTEM OVERRIDE.", 11, False),
        ("", 11, False),
        ("The one thing I most want to know", 12, True),
        (QUESTIONS.get(code, "Does the text read naturally overall?"), 11, False),
        ("", 11, False),
        ("When you are done, reply to the email with this file attached. Thank you.", 11, False),
    ]
    for i, (text, size, bold) in enumerate(lines, start=1):
        c = start.cell(row=i, column=1, value=text)
        c.font = Font(size=size, bold=bold)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if len(text) > 95:
            start.row_dimensions[i].height = 15 * (len(text) // 95 + 1)

    ws = wb.create_sheet("Review")
    cols = [("#", 5), ("Type", 24), ("English", 48), (meta["name"], 48),
            ("Verdict", 12), ("Better version (optional)", 40), ("Notes (optional)", 28), ("id", 14)]
    head_fill = PatternFill("solid", fgColor="1F2328")
    for j, (name, width) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=j, value=name)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = width
    ws.freeze_panes = "A2"

    answer_fill = PatternFill("solid", fgColor="FFF8DB")
    dv = DataValidation(type="list", formula1='"ok,awkward,wrong"', allow_blank=True,
                        showDropDown=False)
    dv.error, dv.errorTitle = "Please pick ok, awkward, or wrong.", "Verdict"
    ws.add_data_validation(dv)

    for n, r in enumerate(rows_for(bank, code), start=1):
        item_id, _kind, category, english, translation = r[:5]
        row = n + 1
        values = [n, CATEGORY_LABEL.get(category, category), english, translation, "", "", "", item_id]
        for j, v in enumerate(values, start=1):
            c = ws.cell(row=row, column=j, value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=row, column=5).fill = answer_fill
        ws.cell(row=row, column=8).font = Font(color="9AA0A6", size=9)
        dv.add(ws.cell(row=row, column=5))
        longest = max(len(english), len(translation))
        ws.row_dimensions[row].height = max(30, 15 * (longest // 52 + 1))

    wb.save(path)
    return path


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("langs", nargs="*", help="language codes")
    ap.add_argument("--tier", choices=["high", "mid", "low"])
    ap.add_argument("--all", action="store_true")
    ap.add_argument("--xlsx", action="store_true",
                    help="also write a formatted Excel sheet for sending to a reviewer")
    args = ap.parse_args()

    bank = load_bank()
    available = list(bank["languages"])
    if args.all:
        codes = available
    elif args.tier:
        codes = [c for c in available if tier_of(c) == args.tier]
    else:
        codes = args.langs
    codes = [c for c in codes if c != "en"]        # nothing to review in the source

    unknown = [c for c in codes if c not in bank["languages"]]
    if unknown:
        print(f"not in the bank: {', '.join(unknown)}")
        return 1
    if not codes:
        print("Nothing selected. Pass language codes, --tier, or --all.")
        print(f"In the bank: {', '.join(c for c in available if c != 'en')}")
        return 1

    for code in codes:
        path = write_sheet(bank, code)
        n = len(rows_for(bank, code))
        print(f"  {bank['languages'][code]['name']:<14} {n:>3} rows -> {path}")
        if args.xlsx:
            print(f"  {'':<14}     + {write_xlsx(bank, code)}")
    print(f"\n{len(codes)} sheet(s) in {OUT_DIR}/")
    print("Send to a speaker of each language. Record results in NATIVE_REVIEW.md")
    print("and set native_reviewed for that language once applied.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
